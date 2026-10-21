import io
import pytest
from app.services.parser import DocumentParser


def test_clean_text():
    dirty = "Software  Engineer\r\n\r\n\r\n\u2022 Responsible for API\t\ufb01les\u201d"
    cleaned = DocumentParser.clean_text(dirty)
    assert "Responsible for API" in cleaned
    assert "files" in cleaned  # Ligature fi normalized
    assert "\r" not in cleaned
    assert "\n\n\n" not in cleaned


def test_parse_txt():
    sample = "Alice Smith\nSoftware Engineer\nPython, Docker, FastAPI"
    raw_bytes = sample.encode("utf-8")
    text, meta = DocumentParser.parse_txt(raw_bytes)
    assert "Alice Smith" in text
    assert meta["file_type"] == "txt"
    assert meta["word_count"] > 0


def test_parse_docx():
    import docx

    doc = docx.Document()
    doc.add_heading("Bob Johnson - Resume", level=1)
    doc.add_paragraph("Senior DevOps Engineer with AWS and Kubernetes expertise.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skill: Terraform"
    table.rows[0].cells[1].text = "Level: Expert"

    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    text, meta = DocumentParser.parse_docx(docx_bytes)
    assert "Bob Johnson" in text
    assert "DevOps Engineer" in text
    assert "Terraform" in text
    assert meta["file_type"] == "docx"


def test_parse_pdf():
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Charlie Davis\nPrincipal Architect\nPython, Distributed Systems")
    pdf_bytes = doc.write()
    doc.close()

    text, meta = DocumentParser.parse_pdf(pdf_bytes)
    assert "Charlie Davis" in text
    assert "Principal Architect" in text
    assert meta["page_count"] == 1
    assert meta["file_type"] == "pdf"


def test_parse_document_dispatcher():
    content = b"Candidate Resume with Python skills"
    text, meta = DocumentParser.parse_document(content, "resume.txt")
    assert "Python skills" in text
    assert meta["file_type"] == "txt"
