import io
import re
from typing import Tuple, Dict, Any, Optional


class DocumentParser:
    """
    Production-grade document parser supporting PDF (PyMuPDF & pdfplumber),
    DOCX (python-docx), and plain text.
    """

    @staticmethod
    def clean_text(text: str) -> str:
        """
        Normalize text, ligatures, bullet points, and excessive whitespace.
        """
        if not text:
            return ""

        # Normalize common unicode ligatures and special quotes
        replacements = {
            "\u201c": '"',
            "\u201d": '"',
            "\u2018": "'",
            "\u2019": "'",
            "\ufb01": "fi",
            "\ufb02": "fl",
            "\u2022": "\n* ",
            "\u25cf": "\n* ",
            "\u25cb": "\n* ",
            "\u25aa": "\n* ",
            "\u2013": "-",
            "\u2014": "--",
            "\t": " ",
            "\r\n": "\n",
            "\r": "\n",
        }
        for orig, rep in replacements.items():
            text = text.replace(orig, rep)

        # Remove null characters and non-printable control codes (except newline)
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)

        # Normalize 3+ newlines to 2 newlines
        text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)

        # Normalize multiple spaces into single space
        text = re.sub(r"[ ]{2,}", " ", text)

        return text.strip()

    @classmethod
    def parse_pdf(cls, file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        """
        Extract text from PDF using PyMuPDF (fitz) with fallback to pdfplumber.
        """
        text_content = []
        page_count = 0
        method_used = "pymupdf"

        try:
            import pymupdf
            doc = pymupdf.open(stream=file_bytes, filetype="pdf")
            page_count = len(doc)
            for page in doc:
                text_content.append(page.get_text("text"))
            raw_text = "\n".join(text_content).strip()
        except Exception:
            raw_text = ""

        # Fallback to pdfplumber if fitz produced empty or failed
        if not raw_text.strip():
            try:
                import pdfplumber
                method_used = "pdfplumber"
                with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                    page_count = len(pdf.pages)
                    for page in pdf.pages:
                        extracted = page.extract_text()
                        if extracted:
                            text_content.append(extracted)
                raw_text = "\n".join(text_content).strip()
            except Exception as e:
                # If neither library is installed or both failed, handle gracefully
                method_used = f"failed: {str(e)}"
                raw_text = ""

        cleaned = cls.clean_text(raw_text)
        metadata = {
            "page_count": page_count,
            "word_count": len(cleaned.split()),
            "character_count": len(cleaned),
            "engine": method_used,
            "file_type": "pdf",
        }
        return cleaned, metadata

    @classmethod
    def parse_docx(cls, file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        """
        Extract text from Microsoft Word (.docx) files.
        """
        text_content = []
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)

            # Extract text from tables if any
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        text_content.append(" | ".join(row_data))

            raw_text = "\n".join(text_content)
        except Exception as e:
            raw_text = ""

        cleaned = cls.clean_text(raw_text)
        metadata = {
            "page_count": 1,  # DOCX has flow layout
            "word_count": len(cleaned.split()),
            "character_count": len(cleaned),
            "engine": "python-docx",
            "file_type": "docx",
        }
        return cleaned, metadata

    @classmethod
    def parse_txt(cls, file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        """
        Extract text from raw TXT or Markdown files with encoding detection.
        """
        raw_text = ""
        for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                raw_text = file_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        cleaned = cls.clean_text(raw_text)
        metadata = {
            "page_count": 1,
            "word_count": len(cleaned.split()),
            "character_count": len(cleaned),
            "engine": "text_decoder",
            "file_type": "txt",
        }
        return cleaned, metadata

    @classmethod
    def parse_document(cls, file_bytes: bytes, filename: str) -> Tuple[str, Dict[str, Any]]:
        """
        Dispatches document parsing based on filename extension.
        """
        lower_name = filename.lower()
        if lower_name.endswith(".pdf"):
            return cls.parse_pdf(file_bytes)
        elif lower_name.endswith(".docx") or lower_name.endswith(".doc"):
            return cls.parse_docx(file_bytes)
        elif lower_name.endswith((".txt", ".md", ".rtf", ".json")):
            return cls.parse_txt(file_bytes)
        else:
            # Default to text decoding attempt
            return cls.parse_txt(file_bytes)
